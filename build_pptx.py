#!/usr/bin/env python3
"""Генератор презентации о Луне в формате PowerPoint (.pptx).

Запуск:
    python3 build_pptx.py
Результат:
    luna.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# --- Палитра (космическая тема) ---
BG_TOP = RGBColor(0x10, 0x18, 0x3A)
BG_BOTTOM = RGBColor(0x05, 0x06, 0x0F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x8A, 0xB4, 0xFF)
ACCENT_SOFT = RGBColor(0xC9, 0xD8, 0xFF)
MUTED = RGBColor(0x9A, 0xA6, 0xC7)
CARD = RGBColor(0x18, 0x20, 0x44)

# 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
BLANK = prs.slide_layouts[6]


def add_gradient_bg(slide):
    """Вертикальный градиентный фон."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    shape.line.fill.background()
    shape.shadow.inherit = False
    fill = shape.fill
    fill.gradient()
    try:
        fill.gradient_angle = 90.0
    except Exception:
        pass
    stops = fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = BG_TOP
    stops[1].position = 1.0
    stops[1].color.rgb = BG_BOTTOM
    # отправить на задний план
    spTree = slide.shapes._spTree
    spTree.remove(shape._element)
    spTree.insert(2, shape._element)
    return shape


def add_moon(slide, left, top, size):
    """Декоративный круг-Луна с радиальным градиентом."""
    moon = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    moon.line.color.rgb = RGBColor(0x6F, 0x80, 0xB5)
    moon.line.width = Pt(0.75)
    moon.shadow.inherit = False
    fill = moon.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    stops[1].position = 1.0
    stops[1].color.rgb = RGBColor(0x8E, 0x9B, 0xC4)
    # радиальный градиент из верхнего-левого угла
    grad = fill._xPr.find(qn('a:gradFill'))
    if grad is not None:
        for child in list(grad):
            if child.tag in (qn('a:lin'), qn('a:path')):
                grad.remove(child)
        path = grad.makeelement(qn('a:path'), {'path': 'circle'})
        fillToRect = path.makeelement(qn('a:fillToRect'),
                                      {'l': '30000', 't': '30000', 'r': '70000', 'b': '70000'})
        path.append(fillToRect)
        grad.append(path)
    return moon


def set_text(tf, runs_spec, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=Pt(6), line_spacing=1.1):
    """runs_spec: список параграфов; каждый параграф — список (text, size, color, bold)."""
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for para in runs_spec:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = space_after
        p.line_spacing = line_spacing
        for (text, size, color, bold) in para:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            r.font.name = "Calibri"


def add_title(slide, text, top=Inches(0.55)):
    box = slide.shapes.add_textbox(Inches(0.9), top, Inches(11.5), Inches(1.1))
    set_text(box.text_frame, [[(text, 40, WHITE, True)]])
    # подчёркивающая линия-акцент
    line = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(0.95), top + Inches(1.0), Inches(1.0), Pt(5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()
    line.shadow.inherit = False
    return box


def add_bullets(slide, items, top=Inches(1.9), left=Inches(0.95), width=Inches(11.4)):
    box = slide.shapes.add_textbox(left, top, width, Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(14)
        p.line_spacing = 1.15
        # маркер
        bullet = p.add_run()
        bullet.text = "●  "
        bullet.font.size = Pt(16)
        bullet.font.color.rgb = ACCENT
        # текст (поддержка выделения жирным через список кортежей)
        parts = item if isinstance(item, list) else [(item, False)]
        for txt, bold in parts:
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(20)
            r.font.bold = bold
            r.font.color.rgb = WHITE if bold else ACCENT_SOFT
            r.font.name = "Calibri"
    return box


def add_fact_card(slide, left, top, w, h, value, label):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.color.rgb = ACCENT
    card.line.width = Pt(1)
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(tf, [
        [(value, 24, WHITE, True)],
        [(label, 13, MUTED, False)],
    ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, space_after=Pt(4))


# ============ СЛАЙД 1: Титульный ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_moon(s, Inches(5.42), Inches(1.1), Inches(2.5))
box = s.shapes.add_textbox(Inches(1.0), Inches(3.9), Inches(11.3), Inches(2.2))
set_text(box.text_frame, [
    [("ЛУНА", 66, WHITE, True)],
    [("Единственный естественный спутник Земли", 24, ACCENT_SOFT, False)],
], align=PP_ALIGN.CENTER, space_after=Pt(10))

# ============ СЛАЙД 2: Что такое Луна ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_title(s, "Что такое Луна?")
add_bullets(s, [
    "Луна — единственный естественный спутник Земли и пятый по величине спутник в Солнечной системе.",
    "Это ближайшее к нам космическое тело и единственное, на котором побывал человек.",
    "Луна влияет на жизнь Земли: вызывает приливы и отливы, стабилизирует наклон земной оси.",
    "Она всегда повёрнута к Земле одной стороной — это называется приливным захватом.",
])

# ============ СЛАЙД 3: Ключевые факты ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_title(s, "Ключевые факты")
facts = [
    ("3 474 км", "Диаметр (≈¼ земного)"),
    ("384 400 км", "Среднее расстояние до Земли"),
    ("27,3 суток", "Период обращения вокруг Земли"),
    ("1/6 g", "Сила тяжести от земной"),
    ("−173…+127 °C", "Перепад температур"),
    ("4,5 млрд лет", "Возраст"),
]
cols, rows = 3, 2
gap = Inches(0.35)
margin = Inches(0.95)
cw = (SLIDE_W - margin * 2 - gap * (cols - 1)) / cols
ch = Inches(1.7)
start_top = Inches(2.1)
for i, (val, lab) in enumerate(facts):
    r, c = divmod(i, cols)
    left = margin + c * (cw + gap)
    top = start_top + r * (ch + gap)
    add_fact_card(s, left, top, cw, ch, val, lab)

# ============ СЛАЙД 4: Происхождение ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_title(s, "Происхождение Луны")
add_bullets(s, [
    [("Гипотеза гигантского столкновения. ", True),
     ("Около 4,5 млрд лет назад в молодую Землю врезалось тело размером с Марс — Тейя.", False)],
    "Из выброшенных обломков и пыли на околоземной орбите сформировалась Луна.",
    "Эта теория объясняет схожий состав пород Земли и Луны, а также большой размер спутника.",
    "Существуют и другие гипотезы (захват, совместное образование), но столкновение — основная.",
])

# ============ СЛАЙД 5: Строение ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_title(s, "Строение Луны")
add_bullets(s, [
    [("Кора", True), (" — внешний твёрдый слой толщиной около 50 км.", False)],
    [("Мантия", True), (" — основная часть объёма, богатая силикатами.", False)],
    [("Ядро", True), (" — небольшое, частично расплавленное, радиусом около 350 км.", False)],
    "Поверхность покрыта реголитом — слоем пыли и обломков от метеоритных ударов.",
    "Тёмные «моря» — застывшая лава, светлые области — древние горные материки.",
])

# ============ СЛАЙД 6: Фазы Луны ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_title(s, "Фазы Луны")
sub = s.shapes.add_textbox(Inches(0.95), Inches(1.75), Inches(11.4), Inches(0.6))
set_text(sub.text_frame, [[("Полный цикл смены фаз — синодический месяц, около 29,5 суток.", 18, MUTED, False)]])
phases = ["Новолуние", "Растущий серп", "Первая четверть",
          "Полнолуние", "Последняя четверть", "Убывающий серп"]
cols = 3
cw = Inches(3.6)
ch = Inches(1.4)
gap = Inches(0.45)
margin = Inches(0.95)
start_top = Inches(2.7)
for i, name in enumerate(phases):
    r, c = divmod(i, cols)
    left = margin + c * (cw + gap)
    top = start_top + r * (ch + gap)
    circ = s.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.9), Inches(0.9))
    circ.fill.solid()
    circ.fill.fore_color.rgb = ACCENT_SOFT
    circ.line.color.rgb = WHITE
    circ.line.width = Pt(0.75)
    circ.shadow.inherit = False
    lbl = s.shapes.add_textbox(left + Inches(1.05), top, cw - Inches(1.05), Inches(0.9))
    set_text(lbl.text_frame, [[(name, 18, ACCENT_SOFT, False)]], anchor=MSO_ANCHOR.MIDDLE)

# ============ СЛАЙД 7: Влияние на Землю ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_title(s, "Влияние на Землю")
add_bullets(s, [
    [("Приливы и отливы. ", True), ("Притяжение Луны поднимает уровень океанов дважды в сутки.", False)],
    [("Стабилизация оси. ", True), ("Луна удерживает наклон земной оси, делая климат устойчивым.", False)],
    [("Замедление вращения. ", True), ("Из-за приливного трения сутки на Земле постепенно удлиняются.", False)],
    [("Удаление. ", True), ("Луна отдаляется от Земли примерно на 3,8 см в год.", False)],
])

# ============ СЛАЙД 8: Исследование (хронология) ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_title(s, "Исследование Луны")
timeline = [
    ("1959", "«Луна-2» — первый аппарат, достигший поверхности Луны."),
    ("1966", "«Луна-9» — первая мягкая посадка."),
    ("1969", "«Аполлон-11» — первые люди на Луне: Армстронг и Олдрин."),
    ("1970", "«Луноход-1» — первый планетоход на другом небесном теле."),
    ("2020-е", "Новая лунная гонка: программа «Артемида», миссии Китая и Индии."),
]
box = s.shapes.add_textbox(Inches(0.95), Inches(2.0), Inches(11.4), Inches(5.0))
tf = box.text_frame
tf.word_wrap = True
first = True
for year, desc in timeline:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    first = False
    p.space_after = Pt(16)
    p.line_spacing = 1.1
    ry = p.add_run()
    ry.text = f"{year}   "
    ry.font.size = Pt(22)
    ry.font.bold = True
    ry.font.color.rgb = ACCENT
    rd = p.add_run()
    rd.text = desc
    rd.font.size = Pt(20)
    rd.font.color.rgb = ACCENT_SOFT

# ============ СЛАЙД 9: Интересные факты ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_title(s, "Интересные факты")
add_bullets(s, [
    "На Луне нет атмосферы, поэтому небо там всегда чёрное, а звук не распространяется.",
    "Следы астронавтов могут сохраняться миллионы лет — их некому стереть.",
    "С Земли видно всегда около 59% поверхности Луны.",
    "Свет от Луны доходит до Земли примерно за 1,3 секунды.",
    "Полное лунное затмение окрашивает Луну в красный цвет — «кровавая Луна».",
])

# ============ СЛАЙД 10: Заключение ============
s = prs.slides.add_slide(BLANK)
add_gradient_bg(s)
add_moon(s, Inches(5.79), Inches(1.3), Inches(1.75))
box = s.shapes.add_textbox(Inches(1.0), Inches(3.6), Inches(11.3), Inches(2.4))
set_text(box.text_frame, [
    [("Спасибо за внимание!", 40, WHITE, True)],
    [("Луна — наш ближайший космический сосед и ключ к пониманию Солнечной системы.", 22, ACCENT_SOFT, False)],
], align=PP_ALIGN.CENTER, space_after=Pt(12))

prs.save("luna.pptx")
print(f"Готово: luna.pptx, слайдов: {len(prs.slides._sldIdLst)}")
